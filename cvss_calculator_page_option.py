# cvss_calculator_page.py
# -------------------------------------------------------------
# CVSS v3.1 / v4.0 계산기 데모 — "CVE/약점에 CVSS가 없을 때 직접 산정" 기능
# 계산 엔진: cvss_calc.py (v3.1 닫힌 수식 + v4.0 FIRST 공식 참조구현 포팅) 그대로 이식
# UI 문구: Weakness_Vulnerability_Template_v2.xlsm의
#          'v3.1_벡터생성기'/'v4.0_벡터생성기' 시트 문구를 그대로 재현
#          (코드 암기 없이 질문에 답하듯 선택하면 벡터/점수가 자동 생성됨)
#
# [설계 결정]
#   - v4.0_벡터생성기 시트 자체는 "미세 보간 미반영, ±0.3~0.5 오차 가능"이라고
#     명시하고 정확한 값은 RunCVSSCalc(=cvss_calc.py) 실행을 권장한다. 이 데모는
#     처음부터 cvss_calc.py의 정확한 엔진으로 계산하므로 그 안내와 동일한 수준의
#     정확도를 UI 단계에서 바로 제공한다.
#   - Weakness_Vulnerability_Template_v2.xlsm을 불러오면, CVSS_V3.1/CVSS_V4 열이
#     비어 있는 행만 골라 보여주고, 계산 결과를 그 행에 바로 반영할 수 있다.
# -------------------------------------------------------------

import os
import io
import re

import streamlit as st
import pandas as pd
import openpyxl

DEFAULT_EXCEL_FILENAME = "Weakness_Vulnerability_Template_v2.xlsm"
SHEET_MAIN = "Weakness_Vulnerability"
HEADER_ROW = 4
DATA_START_ROW = 5

COL_WEAKNESS_ID = 3
COL_TRIGGER_ID = 4
COL_VENDOR = 8
COL_PRODUCT = 9
COL_CWE = 6
COL_CVSS_V4 = 17   # Q
COL_CVSS_V31 = 18  # R


# ======================================================================
# CVSS 계산 엔진 (cvss_calc.py 그대로 이식)
# ======================================================================
_V31_AV = {"N": 0.85, "A": 0.62, "L": 0.55, "P": 0.20}
_V31_AC = {"L": 0.77, "H": 0.44}
_V31_PR_U = {"N": 0.85, "L": 0.62, "H": 0.27}
_V31_PR_C = {"N": 0.85, "L": 0.68, "H": 0.50}
_V31_UI = {"N": 0.85, "R": 0.62}
_V31_CIA = {"H": 0.56, "L": 0.22, "N": 0.00}


def _parse_vector(vector):
    vector = vector.replace("CVSS:3.1/", "").replace("CVSS:4.0/", "")
    parts = [p for p in vector.split("/") if p]
    return {k: v for k, v in (p.split(":") for p in parts)}


def _roundup(x):
    int_input = round(x * 100000)
    if int_input % 10000 == 0:
        return int_input / 100000.0
    return (int_input // 10000 + 1) / 10.0


def cvss31_base_score(vector):
    m = _parse_vector(vector)
    av, ac, pr, ui = m["AV"], m["AC"], m["PR"], m["UI"]
    scope = m["S"]
    c, i, a = m["C"], m["I"], m["A"]

    c_v, i_v, a_v = _V31_CIA[c], _V31_CIA[i], _V31_CIA[a]
    isc_base = 1 - ((1 - c_v) * (1 - i_v) * (1 - a_v))

    impact = 6.42 * isc_base if scope == "U" else 7.52 * (isc_base - 0.029) - 3.25 * ((isc_base - 0.02) ** 15)
    pr_table = _V31_PR_U if scope == "U" else _V31_PR_C
    exploitability = 8.22 * _V31_AV[av] * _V31_AC[ac] * pr_table[pr] * _V31_UI[ui]

    if impact <= 0:
        base = 0.0
    elif scope == "U":
        base = _roundup(min(impact + exploitability, 10))
    else:
        base = _roundup(min(1.08 * (impact + exploitability), 10))
    return round(base, 1)


cvssLookup_global = {
    "000000": 10, "000001": 9.9, "000010": 9.8, "000011": 9.5, "000020": 9.5, "000021": 9.2,
    "000100": 10, "000101": 9.6, "000110": 9.3, "000111": 8.7, "000120": 9.1, "000121": 8.1,
    "000200": 9.3, "000201": 9, "000210": 8.9, "000211": 8, "000220": 8.1, "000221": 6.8,
    "001000": 9.8, "001001": 9.5, "001010": 9.5, "001011": 9.2, "001020": 9, "001021": 8.4,
    "001100": 9.3, "001101": 9.2, "001110": 8.9, "001111": 8.1, "001120": 8.1, "001121": 6.5,
    "001200": 8.8, "001201": 8, "001210": 7.8, "001211": 7, "001220": 6.9, "001221": 4.8,
    "002001": 9.2, "002011": 8.2, "002021": 7.2, "002101": 7.9, "002111": 6.9, "002121": 5,
    "002201": 6.9, "002211": 5.5, "002221": 2.7, "010000": 9.9, "010001": 9.7, "010010": 9.5,
    "010011": 9.2, "010020": 9.2, "010021": 8.5, "010100": 9.5, "010101": 9.1, "010110": 9,
    "010111": 8.3, "010120": 8.4, "010121": 7.1, "010200": 9.2, "010201": 8.1, "010210": 8.2,
    "010211": 7.1, "010220": 7.2, "010221": 5.3, "011000": 9.5, "011001": 9.3, "011010": 9.2,
    "011011": 8.5, "011020": 8.5, "011021": 7.3, "011100": 9.2, "011101": 8.2, "011110": 8,
    "011111": 7.2, "011120": 7, "011121": 5.9, "011200": 8.4, "011201": 7, "011210": 7.1,
    "011211": 5.2, "011220": 5, "011221": 3, "012001": 8.6, "012011": 7.5, "012021": 5.2,
    "012101": 7.1, "012111": 5.2, "012121": 2.9, "012201": 6.3, "012211": 2.9, "012221": 1.7,
    "100000": 9.8, "100001": 9.5, "100010": 9.4, "100011": 8.7, "100020": 9.1, "100021": 8.1,
    "100100": 9.4, "100101": 8.9, "100110": 8.6, "100111": 7.4, "100120": 7.7, "100121": 6.4,
    "100200": 8.7, "100201": 7.5, "100210": 7.4, "100211": 6.3, "100220": 6.3, "100221": 4.9,
    "101000": 9.4, "101001": 8.9, "101010": 8.8, "101011": 7.7, "101020": 7.6, "101021": 6.7,
    "101100": 8.6, "101101": 7.6, "101110": 7.4, "101111": 5.8, "101120": 5.9, "101121": 5,
    "101200": 7.2, "101201": 5.7, "101210": 5.7, "101211": 5.2, "101220": 5.2, "101221": 2.5,
    "102001": 8.3, "102011": 7, "102021": 5.4, "102101": 6.5, "102111": 5.8, "102121": 2.6,
    "102201": 5.3, "102211": 2.1, "102221": 1.3, "110000": 9.5, "110001": 9, "110010": 8.8,
    "110011": 7.6, "110020": 7.6, "110021": 7, "110100": 9, "110101": 7.7, "110110": 7.5,
    "110111": 6.2, "110120": 6.1, "110121": 5.3, "110200": 7.7, "110201": 6.6, "110210": 6.8,
    "110211": 5.9, "110220": 5.2, "110221": 3, "111000": 8.9, "111001": 7.8, "111010": 7.6,
    "111011": 6.7, "111020": 6.2, "111021": 5.8, "111100": 7.4, "111101": 5.9, "111110": 5.7,
    "111111": 5.7, "111120": 4.7, "111121": 2.3, "111200": 6.1, "111201": 5.2, "111210": 5.7,
    "111211": 2.9, "111220": 2.4, "111221": 1.6, "112001": 7.1, "112011": 5.9, "112021": 3,
    "112101": 5.8, "112111": 2.6, "112121": 1.5, "112201": 2.3, "112211": 1.3, "112221": 0.6,
    "200000": 9.3, "200001": 8.7, "200010": 8.6, "200011": 7.2, "200020": 7.5, "200021": 5.8,
    "200100": 8.6, "200101": 7.4, "200110": 7.4, "200111": 6.1, "200120": 5.6, "200121": 3.4,
    "200200": 7, "200201": 5.4, "200210": 5.2, "200211": 4, "200220": 4, "200221": 2.2,
    "201000": 8.5, "201001": 7.5, "201010": 7.4, "201011": 5.5, "201020": 6.2, "201021": 5.1,
    "201100": 7.2, "201101": 5.7, "201110": 5.5, "201111": 4.1, "201120": 4.6, "201121": 1.9,
    "201200": 5.3, "201201": 3.6, "201210": 3.4, "201211": 1.9, "201220": 1.9, "201221": 0.8,
    "202001": 6.4, "202011": 5.1, "202021": 2, "202101": 4.7, "202111": 2.1, "202121": 1.1,
    "202201": 2.4, "202211": 0.9, "202221": 0.4, "210000": 8.8, "210001": 7.5, "210010": 7.3,
    "210011": 5.3, "210020": 6, "210021": 5, "210100": 7.3, "210101": 5.5, "210110": 5.9,
    "210111": 4, "210120": 4.1, "210121": 2, "210200": 5.4, "210201": 4.3, "210210": 4.5,
    "210211": 2.2, "210220": 2, "210221": 1.1, "211000": 7.5, "211001": 5.5, "211010": 5.8,
    "211011": 4.5, "211020": 4, "211021": 2.1, "211100": 6.1, "211101": 5.1, "211110": 4.8,
    "211111": 1.8, "211120": 2, "211121": 0.9, "211200": 4.6, "211201": 1.8, "211210": 1.7,
    "211211": 0.7, "211220": 0.8, "211221": 0.2, "212001": 5.3, "212011": 2.4, "212021": 1.4,
    "212101": 2.4, "212111": 1.2, "212121": 0.5, "212201": 1, "212211": 0.3, "212221": 0.1,
}

maxComposed = {
    "eq1": {0: ["AV:N/PR:N/UI:N/"], 1: ["AV:A/PR:N/UI:N/", "AV:N/PR:L/UI:N/", "AV:N/PR:N/UI:P/"],
            2: ["AV:P/PR:N/UI:N/", "AV:A/PR:L/UI:P/"]},
    "eq2": {0: ["AC:L/AT:N/"], 1: ["AC:H/AT:N/", "AC:L/AT:P/"]},
    "eq3": {
        0: {"0": ["VC:H/VI:H/VA:H/CR:H/IR:H/AR:H/"],
            "1": ["VC:H/VI:H/VA:L/CR:M/IR:M/AR:H/", "VC:H/VI:H/VA:H/CR:M/IR:M/AR:M/"]},
        1: {"0": ["VC:L/VI:H/VA:H/CR:H/IR:H/AR:H/", "VC:H/VI:L/VA:H/CR:H/IR:H/AR:H/"],
            "1": ["VC:L/VI:H/VA:L/CR:H/IR:M/AR:H/", "VC:L/VI:H/VA:H/CR:H/IR:M/AR:M/",
                  "VC:H/VI:L/VA:H/CR:M/IR:H/AR:M/", "VC:H/VI:L/VA:L/CR:M/IR:H/AR:H/",
                  "VC:L/VI:L/VA:H/CR:H/IR:H/AR:M/"]},
        2: {"1": ["VC:L/VI:L/VA:L/CR:H/IR:H/AR:H/"]},
    },
    "eq4": {0: ["SC:H/SI:S/SA:S/"], 1: ["SC:H/SI:H/SA:H/"], 2: ["SC:L/SI:L/SA:L/"]},
    "eq5": {0: ["E:A/"], 1: ["E:P/"], 2: ["E:U/"]},
}

maxSeverity = {
    "eq1": {0: 1, 1: 4, 2: 5}, "eq2": {0: 1, 1: 2},
    "eq3eq6": {0: {0: 7, 1: 6}, 1: {0: 8, 1: 8}, 2: {1: 10}},
    "eq4": {0: 6, 1: 5, 2: 4}, "eq5": {0: 1, 1: 1, 2: 1},
}

_LEVELS = {
    "AV": {"N": 0.0, "A": 0.1, "L": 0.2, "P": 0.3}, "PR": {"N": 0.0, "L": 0.1, "H": 0.2},
    "UI": {"N": 0.0, "P": 0.1, "A": 0.2}, "AC": {"L": 0.0, "H": 0.1}, "AT": {"N": 0.0, "P": 0.1},
    "VC": {"H": 0.0, "L": 0.1, "N": 0.2}, "VI": {"H": 0.0, "L": 0.1, "N": 0.2}, "VA": {"H": 0.0, "L": 0.1, "N": 0.2},
    "SC": {"H": 0.1, "L": 0.2, "N": 0.3}, "SI": {"S": 0.0, "H": 0.1, "L": 0.2, "N": 0.3},
    "SA": {"S": 0.0, "H": 0.1, "L": 0.2, "N": 0.3},
    "CR": {"H": 0.0, "M": 0.1, "L": 0.2}, "IR": {"H": 0.0, "M": 0.1, "L": 0.2}, "AR": {"H": 0.0, "M": 0.1, "L": 0.2},
}


def _m(sel, metric):
    v = sel.get(metric, "X")
    if metric == "E" and v == "X":
        return "A"
    if metric in ("CR", "IR", "AR") and v == "X":
        return "H"
    mkey = "M" + metric
    if mkey in sel and sel[mkey] != "X":
        return sel[mkey]
    return v


def _macro_vector(sel):
    av, pr, ui = _m(sel, "AV"), _m(sel, "PR"), _m(sel, "UI")
    if av == "N" and pr == "N" and ui == "N":
        eq1 = "0"
    elif (av == "N" or pr == "N" or ui == "N") and av != "P":
        eq1 = "1"
    else:
        eq1 = "2"

    ac, at = _m(sel, "AC"), _m(sel, "AT")
    eq2 = "0" if (ac == "L" and at == "N") else "1"

    vc, vi, va = _m(sel, "VC"), _m(sel, "VI"), _m(sel, "VA")
    if vc == "H" and vi == "H":
        eq3 = 0
    elif vc == "H" or vi == "H" or va == "H":
        eq3 = 1
    else:
        eq3 = 2

    sc, si, sa = _m(sel, "SC"), _m(sel, "SI"), _m(sel, "SA")
    msi, msa = _m(sel, "MSI"), _m(sel, "MSA")
    if msi == "S" or msa == "S":
        eq4 = 0
    elif sc == "H" or si == "H" or sa == "H":
        eq4 = 1
    else:
        eq4 = 2

    e = _m(sel, "E")
    eq5 = {"A": 0, "P": 1, "U": 2}[e]

    cr, ir, ar = _m(sel, "CR"), _m(sel, "IR"), _m(sel, "AR")
    eq6 = 0 if ((cr == "H" and vc == "H") or (ir == "H" and vi == "H") or (ar == "H" and va == "H")) else 1

    return f"{eq1}{eq2}{eq3}{eq4}{eq5}{eq6}"


def _extract(metric, s):
    idx = s.index(metric) + len(metric) + 1
    rest = s[idx:]
    return rest[:rest.index("/")] if "/" in rest else rest


def cvss4_base_score(vector):
    sel = _parse_vector(vector)

    if all(_m(sel, k) == "N" for k in ["VC", "VI", "VA", "SC", "SI", "SA"]):
        return 0.0

    mv = _macro_vector(sel)
    value = cvssLookup_global[mv]
    eq1, eq2, eq3, eq4, eq5, eq6 = [int(c) for c in mv]

    def lower(s):
        return cvssLookup_global.get(s)

    eq1_lower = lower(f"{eq1+1}{eq2}{eq3}{eq4}{eq5}{eq6}")
    eq2_lower = lower(f"{eq1}{eq2+1}{eq3}{eq4}{eq5}{eq6}")
    if eq3 == 1 and eq6 == 1:
        eq3eq6_lower = lower(f"{eq1}{eq2}{eq3+1}{eq4}{eq5}{eq6}")
    elif eq3 == 0 and eq6 == 1:
        eq3eq6_lower = lower(f"{eq1}{eq2}{eq3+1}{eq4}{eq5}{eq6}")
    elif eq3 == 1 and eq6 == 0:
        eq3eq6_lower = lower(f"{eq1}{eq2}{eq3}{eq4}{eq5}{eq6+1}")
    elif eq3 == 0 and eq6 == 0:
        left = lower(f"{eq1}{eq2}{eq3}{eq4}{eq5}{eq6+1}")
        right = lower(f"{eq1}{eq2}{eq3+1}{eq4}{eq5}{eq6}")
        eq3eq6_lower = max(left, right) if left is not None and right is not None else (left if left is not None else right)
    else:
        eq3eq6_lower = lower(f"{eq1}{eq2}{eq3+1}{eq4}{eq5}{eq6+1}")
    eq4_lower = lower(f"{eq1}{eq2}{eq3}{eq4+1}{eq5}{eq6}")
    eq5_lower = lower(f"{eq1}{eq2}{eq3}{eq4}{eq5+1}{eq6}")

    eq1_maxes = maxComposed["eq1"][eq1]
    eq2_maxes = maxComposed["eq2"][eq2]
    eq3_eq6_maxes = maxComposed["eq3"][eq3][str(eq6)]
    eq4_maxes = maxComposed["eq4"][eq4]

    max_vectors = [a + b + c + d for a in eq1_maxes for b in eq2_maxes for c in eq3_eq6_maxes for d in eq4_maxes]

    metrics_all = ["AV", "PR", "UI", "AC", "AT", "VC", "VI", "VA", "SC", "SI", "SA", "CR", "IR", "AR"]
    chosen = None
    for mvec in max_vectors:
        dists = {}
        ok = True
        for met in metrics_all:
            lv = _LEVELS[met]
            cur = lv[_m(sel, met)]
            try:
                mx = lv[_extract(met, mvec)]
            except ValueError:
                mx = 0.0
            d = cur - mx
            if d < 0:
                ok = False
                break
            dists[met] = d
        if ok:
            chosen = dists
            break
    if chosen is None:
        chosen = {met: 0.0 for met in metrics_all}

    d_eq1 = chosen["AV"] + chosen["PR"] + chosen["UI"]
    d_eq2 = chosen["AC"] + chosen["AT"]
    d_eq3eq6 = chosen["VC"] + chosen["VI"] + chosen["VA"] + chosen["CR"] + chosen["IR"] + chosen["AR"]

    sc_si_sa_dist = 0.0
    for mvec4 in eq4_maxes:
        try:
            sc_m = _LEVELS["SC"][_extract("SC", mvec4)]
            si_m = _LEVELS["SI"][_extract("SI", mvec4)]
            sa_m = _LEVELS["SA"][_extract("SA", mvec4)]
        except ValueError:
            continue
        d = ((_LEVELS["SC"][_m(sel, "SC")] - sc_m) +
             (_LEVELS["SI"][_m(sel, "SI")] - si_m) +
             (_LEVELS["SA"][_m(sel, "SA")] - sa_m))
        if d >= 0:
            sc_si_sa_dist = d
            break
    d_eq4 = sc_si_sa_dist

    step = 0.1
    max_sev_eq1 = maxSeverity["eq1"][eq1] * step
    max_sev_eq2 = maxSeverity["eq2"][eq2] * step
    max_sev_eq3eq6 = maxSeverity["eq3eq6"][eq3][eq6] * step
    max_sev_eq4 = maxSeverity["eq4"][eq4] * step

    total_n, total_norm = 0, 0.0
    for avail, cur_dist, max_sev in [
        (eq1_lower, d_eq1, max_sev_eq1), (eq2_lower, d_eq2, max_sev_eq2),
        (eq3eq6_lower, d_eq3eq6, max_sev_eq3eq6), (eq4_lower, d_eq4, max_sev_eq4),
        (eq5_lower, 0.0, None),
    ]:
        if avail is None:
            continue
        total_n += 1
        avail_dist = value - avail
        pct = (cur_dist / max_sev) if max_sev else 0.0
        total_norm += avail_dist * pct

    mean_distance = (total_norm / total_n) if total_n else 0.0
    score = value - mean_distance
    score = max(0.0, min(10.0, score))
    return round(score, 1)


# ======================================================================
# "코드 암기 불필요" 질문형 UI 옵션 (v3.1_벡터생성기/v4.0_벡터생성기 시트 문구 그대로)
# 각 튜플: (화면에 보여줄 설명, CVSS 코드)
# ======================================================================
V31_QUESTIONS = [
    ("AV", "공격 벡터 — 어디서 접근해야 하나?", [
        ("인터넷/원격에서 접근 가능 (Network)", "N"),
        ("같은 근거리망(CAN, Wi-Fi 등) 안에서만 (Adjacent)", "A"),
        ("기기에 로그인하거나 케이블 연결 필요 (Local)", "L"),
        ("장비를 물리적으로 뜯거나 만져야 함 (Physical)", "P"),
    ]),
    ("AC", "공격 복잡도 — 조건이 까다로운가?", [
        ("특별한 조건 없이 바로 가능 (Low)", "L"),
        ("특정 타이밍/조건이 맞아야 함 (High)", "H"),
    ]),
    ("PR", "필요 권한 — 사전 권한이 필요한가?", [
        ("권한 필요 없음 (None)", "N"),
        ("일반 사용자 권한 필요 (Low)", "L"),
        ("관리자 권한 필요 (High)", "H"),
    ]),
    ("UI", "사용자 상호작용 — 사용자가 뭔가 해야 하나?", [
        ("사용자 개입 불필요 (None)", "N"),
        ("사용자가 뭔가 해줘야 함 - 클릭, USB꽂기 등 (Required)", "R"),
    ]),
    ("S", "범위 — 다른 영역까지 번지나?", [
        ("이 컴포넌트 안에서 끝남 (Unchanged)", "U"),
        ("다른 보안 영역까지 영향이 번짐 (Changed)", "C"),
    ]),
    ("C", "기밀성 영향", [
        ("영향 없음 (None)", "N"),
        ("일부 정보 유출 (Low)", "L"),
        ("전체 정보 유출 (High)", "H"),
    ]),
    ("I", "무결성 영향", [
        ("영향 없음 (None)", "N"),
        ("일부 데이터 변조 가능 (Low)", "L"),
        ("전체 데이터/펌웨어 임의 변조 가능 (High)", "H"),
    ]),
    ("A", "가용성 영향", [
        ("영향 없음 (None)", "N"),
        ("일부 성능/기능 저하 (Low)", "L"),
        ("완전히 정지/사용불가 (High)", "H"),
    ]),
]

V4_QUESTIONS = [
    ("AV", "공격 벡터 — 어디서 접근해야 하나?", [
        ("인터넷/원격에서 접근 가능 (Network)", "N"),
        ("같은 근거리망(CAN, Wi-Fi 등) 안에서만 (Adjacent)", "A"),
        ("기기에 로그인하거나 케이블 연결 필요 (Local)", "L"),
        ("장비를 물리적으로 뜯거나 만져야 함 (Physical)", "P"),
    ]),
    ("AC", "공격 복잡도 — 방어기법을 우회하는 노력이 필요한가?", [
        ("특별한 우회 없이 바로 가능 (Low)", "L"),
        ("방어기법 우회를 위한 별도 노력 필요 (High)", "H"),
    ]),
    ("AT", "공격 전제조건 — 특정 사전조건이 있어야 하나?", [
        ("전제조건 없음 (None)", "N"),
        ("특정 사전조건(설정, 타이밍 등)이 있어야 함 (Present)", "P"),
    ]),
    ("PR", "필요 권한 — 사전 권한이 필요한가?", [
        ("권한 필요 없음 (None)", "N"),
        ("일반 사용자 권한 필요 (Low)", "L"),
        ("관리자 권한 필요 (High)", "H"),
    ]),
    ("UI", "사용자 상호작용 — 사용자가 뭔가 해야 하나?", [
        ("사용자 개입 불필요 (None)", "N"),
        ("단순히 보기만 함, 별다른 조작 없음 (Passive)", "P"),
        ("적극적으로 실행/설치/클릭 등을 함 (Active)", "A"),
    ]),
    ("VC", "이 컴포넌트 자체의 기밀성 영향", [
        ("영향 없음 (None)", "N"),
        ("일부 정보 유출 (Low)", "L"),
        ("전체 정보 유출 (High)", "H"),
    ]),
    ("VI", "이 컴포넌트 자체의 무결성 영향", [
        ("영향 없음 (None)", "N"),
        ("일부 데이터 변조 가능 (Low)", "L"),
        ("전체 데이터/펌웨어 임의 변조 가능 (High)", "H"),
    ]),
    ("VA", "이 컴포넌트 자체의 가용성 영향", [
        ("영향 없음 (None)", "N"),
        ("일부 성능/기능 저하 (Low)", "L"),
        ("완전히 정지/사용불가 (High)", "H"),
    ]),
    ("SC", "다른 시스템(후속시스템)으로 번지는 기밀성 영향", [
        ("영향 없음 - 이 컴포넌트 안에서 끝남 (None)", "N"),
        ("다른 시스템 일부 정보 유출 (Low)", "L"),
        ("다른 시스템 전체 정보 유출 (High)", "H"),
    ]),
    ("SI", "다른 시스템(후속시스템)으로 번지는 무결성 영향", [
        ("영향 없음 - 이 컴포넌트 안에서 끝남 (None)", "N"),
        ("다른 시스템 일부 데이터 변조 가능 (Low)", "L"),
        ("다른 시스템 전체 데이터 임의 변조 가능 (High)", "H"),
    ]),
    ("SA", "다른 시스템(후속시스템)으로 번지는 가용성 영향", [
        ("영향 없음 - 이 컴포넌트 안에서 끝남 (None)", "N"),
        ("다른 시스템 일부 성능/기능 저하 (Low)", "L"),
        ("다른 시스템 완전 정지/사용불가 (High)", "H"),
    ]),
]


def _render_question_form(questions, key_prefix):
    codes = {}
    for metric, question, options in questions:
        labels = [opt[0] for opt in options]
        picked = st.selectbox(f"**{metric}** — {question}", labels, key=f"{key_prefix}_{metric}")
        codes[metric] = next(code for label, code in options if label == picked)
    return codes


def _codes_to_vector(codes, order):
    return "/".join(f"{m}:{codes[m]}" for m in order)


# ======================================================================
# 엑셀 연동 (선택 사항)
# ======================================================================
def _get_valid_file_path(filename):
    data_folder_path = os.path.join("data", filename)
    if os.path.exists(data_folder_path):
        return data_folder_path
    elif os.path.exists(filename):
        return filename
    return None


def _load_workbook(file_source):
    if isinstance(file_source, str):
        with open(file_source, "rb") as f:
            file_bytes = f.read()
    else:
        file_source.seek(0)
        file_bytes = file_source.read()
    return openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, keep_vba=False)


def _read_missing_cvss_rows(ws):
    rows = []
    r = DATA_START_ROW
    while True:
        wid = ws.cell(row=r, column=COL_WEAKNESS_ID).value
        if wid is None or str(wid).strip() == "":
            probe_empty = all(
                str(ws.cell(row=r + k, column=COL_WEAKNESS_ID).value or "").strip() == "" for k in range(3)
            )
            if probe_empty:
                break
            r += 1
            continue
        v4 = ws.cell(row=r, column=COL_CVSS_V4).value
        v31 = ws.cell(row=r, column=COL_CVSS_V31).value
        if not str(v4 or "").strip() or not str(v31 or "").strip():
            rows.append({
                "row": r, "weakness_id": wid,
                "trigger_id": ws.cell(row=r, column=COL_TRIGGER_ID).value,
                "vendor": ws.cell(row=r, column=COL_VENDOR).value,
                "product": ws.cell(row=r, column=COL_PRODUCT).value,
                "cwe": ws.cell(row=r, column=COL_CWE).value,
                "cvss_v4_now": v4, "cvss_v31_now": v31,
            })
        r += 1
    return rows


# ======================================================================
# Streamlit UI
# ======================================================================
def render():
    st.title("🧮 CVSS v3.1 / v4.0 계산기")
    st.caption("cvss_calc.py(v3.1 닫힌 수식 + v4.0 FIRST 공식 참조구현) 순수 Python 이식 — "
               "CVE/약점에 CVSS가 없을 때, 코드 암기 없이 질문에 답해 정확한 점수를 산정합니다.")

    selected_row = None
    with st.expander("📁 (선택) Weakness_Vulnerability 대장에서 CVSS 없는 항목 불러오기", expanded=False):
        uploaded_file = st.file_uploader(
            "Weakness_Vulnerability_Template 업로드", type=["xlsm", "xlsx"], key="cvss_uploader"
        )
        target_file = uploaded_file if uploaded_file is not None else _get_valid_file_path(DEFAULT_EXCEL_FILENAME)
        if target_file is not None:
            try:
                wb = _load_workbook(target_file)
                if SHEET_MAIN in wb.sheetnames:
                    missing_rows = _read_missing_cvss_rows(wb[SHEET_MAIN])
                    if missing_rows:
                        options = ["(직접 계산만 진행)"] + [
                            f"#{r['row']} {r['weakness_id']} - {r['vendor']} {r['product']}" for r in missing_rows
                        ]
                        picked_label = st.selectbox("CVSS가 비어 있는 항목", options)
                        if picked_label != "(직접 계산만 진행)":
                            idx = options.index(picked_label) - 1
                            selected_row = missing_rows[idx]
                            st.caption(f"CWE: {selected_row['cwe']} · Trigger: {selected_row['trigger_id']}")
                    else:
                        st.info("CVSS가 비어 있는 항목이 없습니다.")
                else:
                    st.warning(f"'{SHEET_MAIN}' 시트를 찾을 수 없습니다.")
            except Exception as e:
                st.error(f"엑셀 로드 오류: {e}")

    tab_v31, tab_v4 = st.tabs(["📐 CVSS v3.1", "📐 CVSS v4.0"])

    with tab_v31:
        st.caption("각 질문에서 취약점 상황에 맞는 설명을 고르면 벡터/점수가 자동 계산됩니다. (v3.1은 닫힌 수식이라 이 값이 곧 정확한 최종값입니다.)")
        codes_31 = _render_question_form(V31_QUESTIONS, "v31")
        vector_31 = _codes_to_vector(codes_31, ["AV", "AC", "PR", "UI", "S", "C", "I", "A"])
        score_31 = cvss31_base_score(vector_31)
        st.markdown("---")
        st.metric("CVSS v3.1 Base Score", score_31)
        st.code(f"CVSS:3.1/{vector_31}", language="text")
        formatted_31 = f"{score_31} (CVSS:3.1/{vector_31})"
        st.text_input("Weakness_Vulnerability 시트 R열(CVSS_V3.1)에 붙여넣을 값", value=formatted_31, key="out_v31")

    with tab_v4:
        st.caption("v4.0은 MacroVector 대표점수에 정밀 보간까지 반영한 정확한 값입니다 (엑셀 벡터생성기 시트의 근사치보다 정밀).")
        codes_4 = _render_question_form(V4_QUESTIONS, "v4")
        vector_4 = _codes_to_vector(codes_4, ["AV", "AC", "AT", "PR", "UI", "VC", "VI", "VA", "SC", "SI", "SA"])
        score_4 = cvss4_base_score(vector_4)
        st.markdown("---")
        st.metric("CVSS v4.0 Base Score", score_4)
        st.code(f"CVSS:4.0/{vector_4}", language="text")
        formatted_4 = f"{score_4} (CVSS:4.0/{vector_4})"
        st.text_input("Weakness_Vulnerability 시트 Q열(CVSS_V4)에 붙여넣을 값", value=formatted_4, key="out_v4")

    if selected_row is not None:
        st.markdown("---")
        st.subheader(f"💾 #{selected_row['row']} {selected_row['weakness_id']}에 반영")
        if st.button("✅ 이 값들을 대장 사본에 반영하고 다운로드 준비"):
            wb2 = _load_workbook(target_file)
            ws2 = wb2[SHEET_MAIN]
            ws2.cell(row=selected_row["row"], column=COL_CVSS_V31).value = formatted_31
            ws2.cell(row=selected_row["row"], column=COL_CVSS_V4).value = formatted_4
            buf = io.BytesIO()
            wb2.save(buf)
            st.session_state["cvss_export"] = buf.getvalue()

        export_bytes = st.session_state.get("cvss_export")
        if export_bytes:
            st.download_button(
                "📥 다운로드: Weakness_Vulnerability_updated.xlsx",
                data=export_bytes,
                file_name="Weakness_Vulnerability_updated.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

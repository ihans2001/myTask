# tara_afr_page.py
# -------------------------------------------------------------
# TARA "4. AFR" 시트 기반 AFR(Attack Feasibility Rating) 자동계산 +
# 공격트리(Attack Tree) 생성 기능의 순수 Python 재구현.
#
# 원본: TARA_template_v1.xlsm 의 VBA 모듈
#   - mAFR_CalcLogic.bas               (Sub CalcAFR_Main)
#   - mStage_4_AFR_AttackTreeGen.bas   (Sub GenerateAttackTree)
#
# [중요] 재구현 시 반영한 단순화/설계 결정 (원본 VBA와 100% 동일하지 않은 부분)
#   1) VBA의 End(xlDown) 블록 끝 탐색은 Excel 특유의 동작(빈 셀을 만나면
#      다음 블록까지 건너뛴 뒤 안전장치로 캡핑됨)이라, 실제로는 거의 항상
#      "다음 헤더행 - N행"으로 캡핑된 값과 동일하게 동작한다. 이를 그대로
#      상수 오프셋으로 대체했다 (AFR 계산: header+33, 공격트리: header+34 -
#      두 원본 모듈이 서로 살짝 다른 상수를 쓰고 있어 각 모듈 고유의 값을 유지).
#   2) 공격트리 생성 시, VBA는 Excel Shape 좌표를 직접 계산해 겹치지 않게
#      배치했지만, Python 버전은 Graphviz의 자동 트리 레이아웃 엔진을 사용한다.
#      노드 종류 판정(ROOT/BOX/ELLIPSE), 색상, AND/OR 표기, 엣지 방향(자식→부모,
#      화살표는 부모 쪽) 등 "의미 있는 로직"은 동일하게 재현했다.
#   3) "그림으로 변환" 옵션(Excel 전용 도형→이미지 변환)은 대상이 없으므로 제외.
#      대신 Graphviz가 화면에 SVG(벡터)로 바로 렌더링해서 보여준다.
#   4) 시나리오 존재 여부/목록은 "4. AFR" 시트 D열(헤더행+1)의 값(예: "TS_0001")
#      을 기준으로 판단한다 (원본 공격트리 모듈은 B열 헤더행 자체를 참조하는데,
#      실제 템플릿에서는 그 위치가 컬럼 라벨 행이라 다소 어긋나 있어 계산 모듈
#      쪽 기준으로 통일했다).
#
# app.py 에서 사용 예:
#   import tara_afr_page
#   ...
#   elif service_mode == "🌳 TARA AFR/공격트리":
#       tara_afr_page.render()
# -------------------------------------------------------------

import os
import io
import streamlit as st
import pandas as pd
import openpyxl
from openpyxl.utils import column_index_from_string
import graphviz

DEFAULT_EXCEL_FILENAME = "TARA_template_v1.xlsm"
SHT_AFR = "4. AFR"

HEADER_ROW_FIRST = 10
BLOCK_SIZE = 35

# AND 결합 시 사용하는 9칸 벡터(G~M) 상의 인덱스 매핑
VEC_COLS = ["G", "H", "I", "J", "K", "L", "M"]  # index 0~6
# (N=Summation -> index 7, O=AFR문자열 -> index 8 은 별도 처리)

METHOD_CONFIG = {
    "HMC4":           {"leaf_sheet": "Ref. Leaf node_HMC",       "combine_cols": ["H", "I", "J", "M"],       "derived": "NONE",    "thresholds": (3, 6, 9, 12)},
    "HEAVENS4":       {"leaf_sheet": "Ref. Leaf node_HEAVENS",   "combine_cols": ["H", "I", "J", "M"],       "derived": "NONE",    "thresholds": (3, 6, 9, 12)},
    "ISO18045_DIM4":  {"leaf_sheet": "Ref. Leaf node_ISO_DIM",   "combine_cols": ["H", "I", "J", "M"],       "derived": "NONE",    "thresholds": (13, 19, 24, 38)},
    "ISO18045_FULL4": {"leaf_sheet": "Ref. Leaf node_ISO_FULL",  "combine_cols": ["G", "H", "I", "J", "M"],  "derived": "NONE",    "thresholds": (13, 19, 24, 57)},
    "FTS_SUM4":       {"leaf_sheet": "Ref. Leaf node_FTS_SUM",   "combine_cols": ["H", "I", "K", "L"],       "derived": "SUM",     "thresholds": (4, 7, 10, 13)},
    "FTS_MULTI4":     {"leaf_sheet": "Ref. Leaf node_FTS_MULTI", "combine_cols": ["H", "I", "K", "L"],       "derived": "PRODUCT", "thresholds": (11, 21, 32, 42)},
}


# =================================================================
# 공용 유틸
# =================================================================
def _get_valid_file_path(filename: str):
    data_folder_path = os.path.join("data", filename)
    if os.path.exists(data_folder_path):
        return data_folder_path
    elif os.path.exists(filename):
        return filename
    return None


def _col(letter):
    return column_index_from_string(letter)


def _cell(ws, row, letter):
    return ws.cell(row=row, column=_col(letter)).value


def _set_cell(ws, row, letter, value):
    ws.cell(row=row, column=_col(letter), value=value)


def _nz_num(v):
    if v is None:
        return 0.0
    if isinstance(v, str):
        v = v.strip()
        if v == "":
            return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def _load_workbook(file_source):
    """무캐싱 실시간 로더 (다른 메뉴들과 동일한 패턴)"""
    if isinstance(file_source, str):
        with open(file_source, "rb") as f:
            file_bytes = f.read()
    else:
        file_source.seek(0)
        file_bytes = file_source.read()
    return openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, keep_vba=False)


def _discover_scenarios(ws):
    """D열(헤더행+1) 값을 기준으로 존재하는 시나리오 블록 목록을 찾는다."""
    scenarios = []
    header_row = HEADER_ROW_FIRST
    max_row = ws.max_row
    while header_row + 1 <= max_row:
        d_val = str(_cell(ws, header_row + 1, "D") or "").strip()
        if not d_val:
            break
        scenarios.append({"header_row": header_row, "top_id": d_val})
        header_row += BLOCK_SIZE
    return scenarios


# =================================================================
# AFR 계산 엔진 (mAFR_CalcLogic.bas 포팅)
# =================================================================
class AFRCalculator:
    def __init__(self, ws, leaf_ws, method_cfg):
        self.ws = ws
        self.leaf_ws = leaf_ws
        self.combine_cols = method_cfg["combine_cols"]
        self.derived = method_cfg["derived"]
        self.thresholds = method_cfg["thresholds"]
        self.processed_ts = {}
        self.processing_set = set()
        self.errors = []
        self.touched_rows = set()   # 화면 표시용: 이번 실행에서 값이 바뀐 행

    # ---- 공통 헬퍼 ----
    @staticmethod
    def _normalize_ts_label(s):
        return str(s).strip().upper().replace("-", "_")

    @staticmethod
    def _is_scenario_id(s):
        s = str(s)
        if s[:3].lower() not in ("ts-", "ts_"):
            return False
        num_part = s[3:]
        if len(num_part) == 0 or "." in num_part:
            return False
        try:
            float(num_part)
        except ValueError:
            return False
        return True

    def _lookup_afr(self, summation):
        b1, b2, b3, b4 = self.thresholds
        if 0 <= summation <= b1:
            return "High"
        elif summation <= b2:
            return "Moderate"
        elif summation <= b3:
            return "Low"
        elif summation <= b4:
            return "Very Low"
        return "ERROR"

    def _find_header_row_by_top_node_id(self, target_id):
        norm_target = self._normalize_ts_label(target_id)
        header_row = HEADER_ROW_FIRST
        max_row = self.ws.max_row
        while header_row + 1 <= max_row:
            d_val = str(_cell(self.ws, header_row + 1, "D") or "").strip()
            if self._normalize_ts_label(d_val) == norm_target:
                return header_row
            header_row += BLOCK_SIZE
        return 0

    def _compute_derived(self, know, equip, expo_t, acc):
        if self.derived == "PRODUCT":
            expert = know * equip
            woo = expo_t * acc
        else:  # SUM
            expert = know + equip
            woo = expo_t + acc
        return expert, woo, expert + woo

    def _set_afr_result(self, r, afr_value):
        _set_cell(self.ws, r, "O", afr_value)
        self.touched_rows.add(r)

    # ---- 리프 노드 채우기 ----
    def _fill_from_leaf_node(self, r, leaf_id):
        last_row_b = self.leaf_ws.max_row
        found_row = None
        for rr in range(5, last_row_b + 1):
            v = _cell(self.leaf_ws, rr, "B")
            if v is not None and str(v).strip().lower() == leaf_id.lower():
                found_row = rr
                break

        if found_row is None:
            for l in ["G", "H", "I", "J", "K", "L", "M", "N"]:
                _set_cell(self.ws, r, l, 0)
            self._set_afr_result(r, "ERROR")
            d_val = str(_cell(self.ws, r, "D") or "")
            self.errors.append(f"[Row {r}] {d_val} : 리프노드 '{leaf_id}'를 리프 시트에서 찾을 수 없음")
            return

        elapsed_t = _nz_num(_cell(self.leaf_ws, found_row, "D"))
        know = _nz_num(_cell(self.leaf_ws, found_row, "F"))
        equip = _nz_num(_cell(self.leaf_ws, found_row, "H"))
        expo_t = _nz_num(_cell(self.leaf_ws, found_row, "K"))
        acc = _nz_num(_cell(self.leaf_ws, found_row, "M"))

        _set_cell(self.ws, r, "G", elapsed_t)
        _set_cell(self.ws, r, "H", know)
        _set_cell(self.ws, r, "I", equip)
        _set_cell(self.ws, r, "K", expo_t)
        _set_cell(self.ws, r, "L", acc)
        self.touched_rows.add(r)

        if self.derived == "NONE":
            _set_cell(self.ws, r, "J", _cell(self.leaf_ws, found_row, "I"))
            _set_cell(self.ws, r, "M", _cell(self.leaf_ws, found_row, "N"))
            _set_cell(self.ws, r, "N", _cell(self.leaf_ws, found_row, "O"))
            self._set_afr_result(r, str(_cell(self.leaf_ws, found_row, "P") or ""))
        else:
            expert, woo, summ = self._compute_derived(know, equip, expo_t, acc)
            _set_cell(self.ws, r, "J", expert)
            _set_cell(self.ws, r, "M", woo)
            _set_cell(self.ws, r, "N", summ)
            self._set_afr_result(r, self._lookup_afr(summ))

    # ---- 다른 시나리오 참조 해석 (재귀 + 순환참조 감지) ----
    def _resolve_scenario_ref(self, scenario_id, caller_row):
        norm_id = self._normalize_ts_label(scenario_id)
        if norm_id not in self.processed_ts:
            h_row = self._find_header_row_by_top_node_id(scenario_id)
            if h_row == 0:
                self.errors.append(f"[Row {caller_row}] {scenario_id} : 참조 시나리오 블록을 시트에서 찾을 수 없음 (0으로 처리)")
                return [0.0] * 8 + ["ERROR"], False
            if norm_id in self.processing_set:
                self.errors.append(f"[Row {caller_row}] {scenario_id} : 순환 참조로 인해 계산할 수 없음 (0으로 처리)")
                return [0.0] * 8 + ["ERROR"], False
            self.process_one_ts(h_row)

        if norm_id in self.processed_ts:
            return self.processed_ts[norm_id], True
        self.errors.append(f"[Row {caller_row}] {scenario_id} : 위협시나리오 계산 실패로 0 처리")
        return [0.0] * 8 + ["ERROR"], False

    def _get_operand_full_vector(self, op_id, d_start, d_last_row, caller_row):
        for rr in range(d_start, d_last_row + 1):
            if str(_cell(self.ws, rr, "D") or "").strip() == op_id:
                vec = [_nz_num(_cell(self.ws, rr, l)) for l in VEC_COLS]
                vec.append(_nz_num(_cell(self.ws, rr, "N")))
                vec.append(str(_cell(self.ws, rr, "O") or ""))
                return vec, True

        if self._is_scenario_id(op_id):
            return self._resolve_scenario_ref(op_id, caller_row)

        return [0.0] * 9, False

    # ---- 조합 노드(AND/OR) 계산 ----
    def _calc_combination(self, r, d_start, d_last_row):
        ts_label = str(_cell(self.ws, r, "D") or "")
        formula = str(_cell(self.ws, r, "F") or "").strip()

        if "||" in formula:
            op_mode, parts = "OR", formula.split("||")
        elif "&" in formula:
            op_mode, parts = "AND", formula.split("&")
        else:
            op_mode, parts = "AND", [formula]

        n_combine = len(self.combine_cols)
        max_vals = [-1.0] * n_combine
        best_sum, best_vec, is_first = -1.0, None, True

        for part in parts:
            op_id = part.strip()
            if not op_id:
                continue
            vec9, ok = self._get_operand_full_vector(op_id, d_start, d_last_row, r)
            if not ok and not self._is_scenario_id(op_id):
                self.errors.append(f"[Row {r}] {ts_label} : Formula 참조 '{op_id}'를 같은 TS 블록에서 찾을 수 없음 (0으로 처리)")

            for ci, colletter in enumerate(self.combine_cols):
                vidx = VEC_COLS.index(colletter)
                v = float(vec9[vidx])
                if v > max_vals[ci]:
                    max_vals[ci] = v

            cur_sum = float(vec9[7])
            if is_first or cur_sum < best_sum:
                best_sum, best_vec, is_first = cur_sum, vec9, False

        self.touched_rows.add(r)

        if op_mode == "OR":
            for i, l in enumerate(VEC_COLS):
                _set_cell(self.ws, r, l, best_vec[i])
            _set_cell(self.ws, r, "N", best_vec[7])
            self._set_afr_result(r, best_vec[8])
        else:
            for l in VEC_COLS:
                _set_cell(self.ws, r, l, "")
            for ci, colletter in enumerate(self.combine_cols):
                _set_cell(self.ws, r, colletter, max_vals[ci])

            if self.derived == "NONE":
                final_sum = sum(max_vals)
            else:
                expert, woo, final_sum = self._compute_derived(*max_vals)
                _set_cell(self.ws, r, "J", expert)
                _set_cell(self.ws, r, "M", woo)

            _set_cell(self.ws, r, "N", final_sum)
            self._set_afr_result(r, self._lookup_afr(final_sum))

    # ---- 블록(TS) 전체 처리 ----
    def process_one_ts(self, header_row):
        d_start = header_row + 1
        ts_label = str(_cell(self.ws, d_start, "D") or "").strip()
        norm_label = self._normalize_ts_label(ts_label)

        if norm_label:
            if norm_label in self.processed_ts:
                return
            if norm_label in self.processing_set:
                self.errors.append(f"[Row {header_row}] {ts_label} : 순환 참조 감지")
                return
            self.processing_set.add(norm_label)

        # 블록 끝 행: 다음 헤더행 - 2 (VBA의 End(xlDown)+안전장치를 상수로 대체 - 모듈 상단 주석 참고)
        next_header_row = header_row + BLOCK_SIZE
        d_last_row = min(next_header_row - 2, self.ws.max_row)

        for rr in range(d_last_row, d_start - 1, -1):
            d_val = str(_cell(self.ws, rr, "D") or "").strip()
            f_val = str(_cell(self.ws, rr, "F") or "").strip()

            if not d_val:
                continue
            elif d_val[:3].lower() == "ts_":
                self._calc_combination(rr, d_start, d_last_row)
            elif d_val[:2].lower() == "t_" and d_val == f_val:
                self._fill_from_leaf_node(rr, d_val)
            else:
                self._set_afr_result(rr, "ERROR")
                self.errors.append(f"[Row {rr}] {ts_label} : D열 값 '{d_val}'을 T_(리프)/TS_(조합) 규칙으로 인식할 수 없음")

        if norm_label:
            vec = [_nz_num(_cell(self.ws, d_start, l)) for l in VEC_COLS]
            vec.append(_nz_num(_cell(self.ws, d_start, "N")))
            vec.append(str(_cell(self.ws, d_start, "O") or ""))
            self.processed_ts[norm_label] = vec
            self.processing_set.discard(norm_label)


# =================================================================
# 공격트리 생성 (mStage_4_AFR_AttackTreeGen.bas 포팅, Graphviz 렌더링)
# =================================================================
class AttackTreeBuilder:
    def __init__(self, ws):
        self.ws = ws
        self.id_dict = {}
        self.referenced = set()
        self.children_cache = {}
        self.formula_cache = {}

    @staticmethod
    def _get_operator(s):
        if "||" in s:
            return "||"
        elif "&" in s:
            return "&"
        return ""

    @staticmethod
    def _is_threat_scenario_ref(id_):
        if len(id_) <= 3 or id_[:3] != "TS_":
            return False
        try:
            float(id_[3:])
            return True
        except ValueError:
            return False

    def build(self, start_row, last_row):
        self.id_dict = {}
        for r in range(start_row, last_row + 1):
            id_ = str(_cell(self.ws, r, "D") or "").strip()
            if id_ and id_ not in self.id_dict:
                self.id_dict[id_] = r

        self.referenced = set()
        for r in range(start_row, last_row + 1):
            id_ = str(_cell(self.ws, r, "D") or "").strip()
            formula = str(_cell(self.ws, r, "F") or "").strip()
            if formula and formula != id_:
                op = self._get_operator(formula)
                if op:
                    for t in formula.split(op):
                        t = t.strip()
                        if t:
                            self.referenced.add(t)
                else:
                    self.referenced.add(formula)

        self.children_cache = {}
        self.formula_cache = {}

    def get_formula(self, id_):
        if id_ in self.formula_cache:
            return self.formula_cache[id_]
        formula = ""
        if id_ in self.id_dict:
            formula = str(_cell(self.ws, self.id_dict[id_], "F") or "").strip()
        self.formula_cache[id_] = formula
        return formula

    def get_children_ids(self, id_):
        if id_ in self.children_cache:
            return self.children_cache[id_]
        formula = self.get_formula(id_)
        result = []
        if formula and formula != id_:
            op = self._get_operator(formula)
            if op:
                result = [t.strip() for t in formula.split(op) if t.strip()]
            else:
                result = [formula]
        self.children_cache[id_] = result
        return result

    def get_shape_kind(self, id_):
        if len(self.get_children_ids(id_)) == 0:
            return "ELLIPSE"
        elif id_ in self.referenced:
            return "BOX"
        return "ROOT"

    def node_style(self, id_):
        formula = self.get_formula(id_)
        op = self._get_operator(formula)
        kind = self.get_shape_kind(id_)
        label = id_
        if op == "&":
            label = f"<AND> {id_}"
        elif op == "||":
            label = f"<OR> {id_}"

        if kind == "ROOT":
            shape, style = "box", "rounded,filled"
            fill, line = "#C6EFCE", "#63A579"
        elif kind == "BOX":
            shape, style = "box", "filled"
            if op == "&":
                fill, line = "#FCD5E3", "#D980A4"
            elif op == "":
                fill, line = "#D6E8FF", "#7AA8E0"
            else:  # "||"
                fill, line = "#CCFFFF", "#00B0C4"
        else:  # ELLIPSE
            shape, style = "ellipse", "filled"
            fill, line = "#FFFFCC", "#D9D997"

        if self._is_threat_scenario_ref(id_):
            fill, line = "#C6EFCE", "#63A579"

        return label, shape, style, fill, line

    def to_graphviz(self, graph_name="attack_tree"):
        dot = graphviz.Digraph(graph_name)
        dot.attr(rankdir="BT")   # Root 위, 리프 아래 (원본과 동일한 시각적 방향)
        dot.attr("node", fontname="Malgun Gothic", fontsize="10")
        dot.attr("edge", color="#808080", arrowhead="normal")

        visited = set()
        roots = [id_ for id_ in self.id_dict if self.get_shape_kind(id_) == "ROOT"]

        def add_subtree(id_):
            if id_ in visited:
                return
            visited.add(id_)
            label, shape, style, fill, line = self.node_style(id_)
            dot.node(id_, label=label, shape=shape, style=style,
                      fillcolor=fill, color=line)
            for child in self.get_children_ids(id_):
                add_subtree(child)
                dot.edge(child, id_)   # 자식 -> 부모 (화살표는 부모 쪽)

        for rid in roots:
            add_subtree(rid)

        return dot, len(roots)


# =================================================================
# Streamlit UI
# =================================================================
def _sidebar_load_workbook():
    with st.expander("📁 TARA 데이터 소스", expanded=False):
        uploaded_file = st.file_uploader(
            "또는 보유하신 TARA 템플릿(.xlsm/.xlsx) 업로드", type=["xlsm", "xlsx"], key="tara_uploader"
        )

    if uploaded_file is not None:
        target_file = uploaded_file
        display_name = uploaded_file.name
        file_valid = True
        current_mtime = getattr(uploaded_file, "size", 0)
    else:
        target_file = _get_valid_file_path(DEFAULT_EXCEL_FILENAME)
        display_name = target_file if target_file else DEFAULT_EXCEL_FILENAME
        file_valid = target_file is not None and os.path.exists(target_file)
        current_mtime = os.path.getmtime(target_file) if file_valid else 0

    if "tara_last_mtime" not in st.session_state:
        st.session_state["tara_last_mtime"] = current_mtime
    elif st.session_state["tara_last_mtime"] != current_mtime:
        st.session_state["tara_last_mtime"] = current_mtime
        st.rerun()

    if not file_valid:
        st.warning(
            f"경고: `{DEFAULT_EXCEL_FILENAME}` 파일이 `data/` 폴더 또는 메인 폴더에 존재하지 않습니다. "
            "사이드바에서 TARA 템플릿 파일을 업로드해 주세요."
        )
        return None, None

    st.success(f"현재 연동 파일: **{display_name}**")
    try:
        wb = _load_workbook(target_file)
    except Exception as e:
        st.error(f"엑셀 파일 로드 중 오류 발생: {e}")
        return None, None

    if SHT_AFR not in wb.sheetnames:
        st.error(f"'{SHT_AFR}' 시트를 찾을 수 없습니다. (시트 목록: {wb.sheetnames})")
        return None, None

    return wb, display_name


def _render_afr_tab(wb):
    ws = wb[SHT_AFR]

    default_method = str(_cell(ws, 1, "C") or "").strip().upper()
    if default_method not in METHOD_CONFIG:
        default_method = "HMC4"

    c1, c2 = st.columns([2, 3])
    with c1:
        method_name = st.selectbox(
            "🧮 AFR 방법론 선택 (시트 C1과 동일한 역할)",
            list(METHOD_CONFIG.keys()),
            index=list(METHOD_CONFIG.keys()).index(default_method),
        )

    method_cfg = METHOD_CONFIG[method_name]
    leaf_sheet_name = method_cfg["leaf_sheet"]

    if leaf_sheet_name not in wb.sheetnames:
        st.error(f"리프 참조 시트 '{leaf_sheet_name}'를 찾을 수 없습니다. (시트 목록: {wb.sheetnames})")
        return

    scenarios = _discover_scenarios(ws)
    if not scenarios:
        st.info(f"'{SHT_AFR}' 시트의 {HEADER_ROW_FIRST + 1}행부터 위협 시나리오 데이터를 찾지 못했습니다.")
        return

    with c2:
        calc_scope = st.radio(
            "계산 범위",
            ["전체 위협시나리오", "선택한 시나리오만"],
            horizontal=True,
        )

    target_scenario = None
    if calc_scope == "선택한 시나리오만":
        target_scenario = st.selectbox(
            "위협시나리오 선택",
            [s["top_id"] for s in scenarios],
        )

    if st.button("🚀 AFR 계산 실행", type="primary"):
        leaf_ws = wb[leaf_sheet_name]
        calc = AFRCalculator(ws, leaf_ws, method_cfg)

        if calc_scope == "전체 위협시나리오":
            for s in scenarios:
                calc.process_one_ts(s["header_row"])
        else:
            target = next(s for s in scenarios if s["top_id"] == target_scenario)
            calc.process_one_ts(target["header_row"])

        st.session_state["tara_afr_calc_result"] = {
            "method_name": method_name,
            "touched_rows": sorted(calc.touched_rows),
            "errors": calc.errors,
        }

    result = st.session_state.get("tara_afr_calc_result")
    if result and result["method_name"] == method_name:
        st.markdown("---")
        touched_rows = result["touched_rows"]
        errors = result["errors"]

        if errors:
            st.error(f"⚠️ 계산 중 {len(errors)}건의 오류가 발생했습니다.")
            with st.expander("오류 목록 보기"):
                for e in errors:
                    st.write(f"- {e}")
        else:
            st.success(f"AFR 계산이 완료되었습니다. (방법론: {method_name}, 오류 없음)")

        if touched_rows:
            def _disp(v):
                return None if v == "" else v

            rows_data = []
            for r in touched_rows:
                rows_data.append({
                    "행": r,
                    "T_ID": _cell(ws, r, "D"),
                    "Formula": _cell(ws, r, "F"),
                    "ElapsedT": _disp(_cell(ws, r, "G")),
                    "Knowledge": _disp(_cell(ws, r, "H")),
                    "Equipment": _disp(_cell(ws, r, "I")),
                    "Expertise": _disp(_cell(ws, r, "J")),
                    "ExposureT": _disp(_cell(ws, r, "K")),
                    "Accessibility": _disp(_cell(ws, r, "L")),
                    "WoO": _disp(_cell(ws, r, "M")),
                    "Summation": _disp(_cell(ws, r, "N")),
                    "AFR": _cell(ws, r, "O"),
                })
            df_result = pd.DataFrame(rows_data)
            for c in ["ElapsedT", "Knowledge", "Equipment", "Expertise", "ExposureT", "Accessibility", "WoO", "Summation"]:
                df_result[c] = pd.to_numeric(df_result[c], errors="coerce")

            def _highlight_error(row):
                return ["background-color: #FFA500" if row["AFR"] == "ERROR" else "" for _ in row]

            st.subheader("📋 계산 결과")
            st.dataframe(df_result.style.apply(_highlight_error, axis=1), use_container_width=True)
            st.caption(
                "💡 원본 엑셀 파일에 포함된 이미지 등으로 인해 '재계산 결과를 원본 형식 그대로 "
                "다시 저장'하는 기능은 안정성 문제로 제공하지 않습니다. 위 표를 CSV로 내보내거나, "
                "필요한 값을 직접 엑셀에 옮겨 반영해 주세요."
            )
            st.download_button(
                "📥 계산 결과 CSV 다운로드",
                data=df_result.to_csv(index=False).encode("utf-8-sig"),
                file_name=f"AFR_result_{method_name}.csv",
                mime="text/csv",
            )


def _render_tree_tab(wb):
    ws = wb[SHT_AFR]
    scenarios = _discover_scenarios(ws)
    if not scenarios:
        st.info(f"'{SHT_AFR}' 시트의 {HEADER_ROW_FIRST + 1}행부터 위협 시나리오 데이터를 찾지 못했습니다.")
        return

    gen_mode = st.radio("생성 방식", ["일괄생성 (전체)", "선택생성 (시나리오 1개)"], horizontal=True)

    targets = scenarios
    if gen_mode == "선택생성 (시나리오 1개)":
        chosen = st.selectbox("공격트리를 생성할 위협시나리오 선택", [s["top_id"] for s in scenarios])
        targets = [s for s in scenarios if s["top_id"] == chosen]

    st.caption(
        "🟩 녹색 둥근사각형 = 최상위(Root) · 🟦 파랑 사각형 = 단일 자식(연산자 없음) · "
        "🩷 분홍 사각형 = AND 결합 · 🩵 시안 사각형 = OR 결합 · 🟨 노랑 타원 = 리프(말단) · "
        "🟩 녹색(도형 무관) = 다른 위협시나리오 참조"
    )

    if st.button("🌳 공격트리 생성", type="primary"):
        builder = AttackTreeBuilder(ws)
        graphs = []
        for s in targets:
            start_row = s["header_row"] + 1
            last_row = s["header_row"] + BLOCK_SIZE - 1
            builder.build(start_row, last_row)
            dot, n_roots = builder.to_graphviz(graph_name=f"tree_{s['top_id']}")
            graphs.append((s["top_id"], dot, n_roots))
        st.session_state["tara_tree_graphs"] = graphs

    graphs = st.session_state.get("tara_tree_graphs")
    if graphs:
        st.markdown("---")
        for scenario_id, dot, n_roots in graphs:
            st.subheader(f"📌 {scenario_id}")
            if n_roots == 0:
                st.warning("최상위(Root) 노드를 찾지 못했습니다.")
                continue
            st.graphviz_chart(dot, use_container_width=True)


def render():
    wb, _display_name = _sidebar_load_workbook()
    if wb is None:
        return

    st.title("🌳 TARA AFR 계산 · 공격트리 생성")
    st.caption("mAFR_CalcLogic.bas / mStage_4_AFR_AttackTreeGen.bas 순수 Python 재구현")

    tab1, tab2 = st.tabs(["🧮 AFR 계산", "🌳 공격트리 생성"])
    with tab1:
        _render_afr_tab(wb)
    with tab2:
        _render_tree_tab(wb)

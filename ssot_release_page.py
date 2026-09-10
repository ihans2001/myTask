# ssot_release_page.py
# -------------------------------------------------------------
# SSoT 기반 릴리즈 준비도 점검 + Draft 릴리즈노트 자동 작성 데모
# 원본: SSoT_Template_v3.xlsm ('인스턴스_통합' / 'SUP9_이슈현황' / '스냅샷_보관' /
#       '릴리즈준비점검' 시트) + modGenerateReleaseNote.bas / finalize_release_hash*.py
#       (문서번호 019, 027 인포그래픽에서 다룬 내용의 실제 동작 버전)
#
# [설계 결정]
#   1) '릴리즈준비점검' 시트의 지표(CR 반영률/CVE 잔존/미종결 Critical 이슈 등)는
#      매크로가 baseline_id 선택 시 그때그때 다시 계산해 값으로 써넣는 방식이라,
#      파일에 저장된 캐시 값은 마지막 실행 시점 것일 뿐 최신이 아닐 수 있다.
#      이 데모는 baseline을 선택할 때마다 '인스턴스_통합'/'SUP9_이슈현황'/
#      '스냅샷_보관'에서 매번 새로 계산한다(캐시된 open_critical_issue_count
#      열도 신뢰하지 않고, SUP9_이슈현황을 직접 대조해 다시 센다).
#   2) 실제 빌드 산출물이 없는 데모 환경이므로, "배포 시점 스냅샷 해시"는 진짜
#      빌드 zip 대신 이번 선택된 baseline 행들의 체크섬을 이어붙여 계산한
#      대표 해시로 대체했다 — 화면에 "데모용 해시"임을 명확히 표시한다.
#   3) Word/PDF 변환은 하지 않고 Markdown(.md)까지만 생성한다(문서번호 027에서
#      다룬 Pandoc/LibreOffice 변환은 로컬 실행파일 의존성이 있어 이 데모의
#      범위 밖으로 뒀다).
# -------------------------------------------------------------

import os
import io
import hashlib
from datetime import date

import streamlit as st
import pandas as pd
import openpyxl

DEFAULT_EXCEL_FILENAME = "SSoT_Template_v3.xlsm"
SHEET_INSTANCE = "인스턴스_통합"
SHEET_SUP9 = "SUP9_이슈현황"
SHEET_SNAPSHOT = "스냅샷_보관"

INSTANCE_HEADER_ROW = 4
INSTANCE_START_ROW = 5
SUP9_HEADER_ROW = 4
SUP9_START_ROW = 5
SNAPSHOT_HEADER_ROW = 4
SNAPSHOT_START_ROW = 5

# 인스턴스_통합 컬럼 (1-based)
COL = {
    "sw_id": 1, "controller_id": 2, "vehicle": 3, "phase": 4, "baseline_id": 5,
    "baseline_status": 6, "version": 7, "cm_reg_id": 8, "cr_id": 9, "ccb_id": 10,
    "reflection_status": 11, "reflection_action": 12, "linked_issue_id": 13,
    "checksum": 14, "detection_source": 15, "last_reviewed_at": 16,
    "safety_relevant_flag": 17, "swe1_ref_id": 18, "vuln_status": 19,
}


def _get_valid_file_path(filename: str):
    data_folder_path = os.path.join("data", filename)
    if os.path.exists(data_folder_path):
        return data_folder_path
    elif os.path.exists(filename):
        return filename
    return None


def _load_workbook(file_source):
    if isinstance(file_source, str):
        with open(file_source, "rb") as f:
            file_bytes = f.read()
    else:
        file_source.seek(0)
        file_bytes = file_source.read()
    return openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, keep_vba=False)


def _cell(ws, row, col):
    return ws.cell(row=row, column=col).value


def _sval(v):
    return "" if v is None else str(v).strip()


# ---------------------------------------------------------------
# 데이터 읽기
# ---------------------------------------------------------------
def _read_instance_rows(ws):
    rows = []
    r = INSTANCE_START_ROW
    while True:
        sw_id = _cell(ws, r, COL["sw_id"])
        if sw_id is None or _sval(sw_id) == "":
            # 중간에 빈 행이 있을 수 있으니 몇 행 더 확인 후 중단
            probe_empty = all(
                _sval(_cell(ws, r + k, COL["sw_id"])) == "" for k in range(0, 3)
            )
            if probe_empty:
                break
            r += 1
            continue
        rows.append({key: _cell(ws, r, idx) for key, idx in COL.items()})
        r += 1
    return rows


def _discover_baselines(instance_rows):
    seen = {}
    for row in instance_rows:
        bid = _sval(row["baseline_id"])
        if not bid:
            continue
        # baseline_id 열에 "BL-A→BL-B→BL-C" 형태(소급 승인 이력)가 들어있는 경우
        # 가장 마지막(최신) baseline만 대표값으로 취급한다.
        last_bid = bid.split("→")[-1].strip()
        seen.setdefault(last_bid, {"vehicle": row["vehicle"], "controller_id": row["controller_id"]})
    return seen


def _read_sup9(ws):
    rows = []
    r = SUP9_START_ROW
    while True:
        issue_id = _cell(ws, r, 1)
        if issue_id is None or _sval(issue_id) == "":
            break
        rows.append({
            "issue_id": _sval(issue_id),
            "severity": _sval(_cell(ws, r, 2)),
            "status": _sval(_cell(ws, r, 3)),
            "related_sw_id": _sval(_cell(ws, r, 4)),
            "updated_at": _sval(_cell(ws, r, 5)),
        })
        r += 1
    return rows


def _read_snapshot_baselines(ws):
    ids = set()
    r = SNAPSHOT_START_ROW
    while True:
        bid = _cell(ws, r, 1)
        if bid is None or _sval(bid) == "":
            break
        ids.add(_sval(bid))
        r += 1
    return ids


# ---------------------------------------------------------------
# 릴리즈 준비도 점검 (판정 로직)
# ---------------------------------------------------------------
def compute_readiness(instance_rows, sup9_rows, snapshot_baselines, baseline_id):
    target_rows = [
        row for row in instance_rows
        if _sval(row["baseline_id"]).split("→")[-1].strip() == baseline_id
    ]

    # 1) CR 반영률 : CR_id가 있는 행 중 reflection_status=='반영완료' 비율
    cr_rows = [row for row in target_rows if _sval(row["cr_id"])]
    reflected = [row for row in cr_rows if _sval(row["reflection_status"]) == "반영완료"]
    unreflected = [row for row in cr_rows if _sval(row["reflection_status"]) != "반영완료"]
    cr_rate = (len(reflected) / len(cr_rows)) if cr_rows else 1.0

    # 2) CVE 미조치/조치중 건수
    cve_pending = [row for row in target_rows if _sval(row["vuln_status"]) in ("미조치", "조치중")]

    # 3) 미종결 Critical 이슈 건수 (SUP9_이슈현황을 직접 대조 - 캐시 열 신뢰 안 함)
    sw_ids = {_sval(row["sw_id"]) for row in target_rows}
    critical_open = [
        issue for issue in sup9_rows
        if issue["severity"] == "Critical" and issue["status"] != "Closed" and issue["related_sw_id"] in sw_ids
    ]

    # 4) 스냅샷 기록 존재 여부
    snapshot_exists = baseline_id in snapshot_baselines

    reasons = []
    if unreflected:
        reasons.append(f"CR 미반영 {len(unreflected)}건 존재 ({', '.join(_sval(r['cr_id']) for r in unreflected)})")
    if cve_pending:
        reasons.append(f"CVE 미조치/조치중 {len(cve_pending)}건 잔존 ({', '.join(_sval(r['sw_id']) for r in cve_pending)})")
    if critical_open:
        reasons.append(f"미종결 Critical 이슈 {len(critical_open)}건 ({', '.join(i['issue_id'] for i in critical_open)})")
    if not snapshot_exists:
        reasons.append("스냅샷_보관 시트 내 해당 baseline_id 미존재")

    verdict = "GO" if not reasons else "NO-GO"

    return {
        "target_rows": target_rows,
        "cr_rate": cr_rate,
        "cr_total": len(cr_rows),
        "cr_reflected": len(reflected),
        "cve_pending": cve_pending,
        "critical_open": critical_open,
        "snapshot_exists": snapshot_exists,
        "verdict": verdict,
        "reasons": reasons,
    }


# ---------------------------------------------------------------
# Draft 릴리즈노트 생성 (★★ 미비 데이터 경고 포함)
# ---------------------------------------------------------------
def generate_draft_release_note(baseline_id, meta, readiness, report_date):
    rows = readiness["target_rows"]
    lines = []
    lines.append("▶ SOFTWARE RELEASE NOTE (DRAFT)")
    lines.append(f"Target Baseline ID: {baseline_id}")
    lines.append(f"Release Date: {report_date}")
    lines.append("Document Version: DRAFT (자동 생성, 연구원 검토·확정 전)")
    lines.append("-" * 70)
    lines.append("")
    lines.append("1. 릴리즈 개요 (Release Overview)")
    lines.append(f"   - 제어기 (Controller ID): {meta.get('controller_id') or '★★ 미확인'}")
    lines.append(f"   - 적용 차종 (Vehicle Model): {meta.get('vehicle') or '★★ 미확인'}")
    lines.append(f"   - 릴리즈 승인 게이트 (Gate): {readiness['verdict']}")
    if readiness["reasons"]:
        lines.append(f"      ㄴ *판정 근거*: {'; '.join(readiness['reasons'])}")
    else:
        lines.append("      ㄴ *판정 근거*: 모든 점검 항목 충족")
    lines.append("")
    lines.append("2. 구성 소프트웨어 컴포넌트 (SWCs)")
    for row in rows:
        version = _sval(row["version"]) or "★★ 버전 미기재"
        lines.append(f"   ㄴ {_sval(row['sw_id'])} (v{version})")
    lines.append("")
    lines.append("3. 주요 변경 사항 (Change Delta - SUP.10)")
    cr_rows = [row for row in rows if _sval(row["cr_id"])]
    if not cr_rows:
        lines.append("   - 이번 baseline에 연결된 CR 없음")
    for row in cr_rows:
        cr_id = _sval(row["cr_id"])
        ccb_id = _sval(row["ccb_id"]) or "★★ CCB 승인번호 없음"
        status = _sval(row["reflection_status"]) or "★★ 반영상태 미기재"
        lines.append(f"   [{cr_id}] {_sval(row['sw_id'])}")
        lines.append(f"      ㄴ CCB 승인번호: {ccb_id} (반영상태: {status})")
        if _sval(row["reflection_action"]):
            lines.append(f"      ㄴ 비고: {_sval(row['reflection_action'])}")
    lines.append("")
    lines.append("4. 알려진 제약사항 및 미해결 이슈 (Known Issues)")
    if readiness["cve_pending"]:
        for row in readiness["cve_pending"]:
            lines.append(f"   - [{_sval(row['sw_id'])}] 취약점 상태: {_sval(row['vuln_status'])}")
    else:
        lines.append("   - 미조치/조치중 취약점 없음")
    if readiness["critical_open"]:
        for issue in readiness["critical_open"]:
            lines.append(f"   - [{issue['issue_id']}] Critical, {issue['status']} (관련 SW: {issue['related_sw_id']})")
    lines.append("")
    lines.append("5. 무결성 검증 및 빌드 체크섬 (Integrity & SHA-256 Checksum)")
    for row in rows:
        checksum = _sval(row["checksum"]) or "★★ 체크섬 미기재"
        lines.append(f"   - {_sval(row['sw_id'])}: `{checksum}`")
    lines.append("   - 10년 보관 스냅샷 파일 해시: `<릴리즈 확정 시 자동 주입 예정>`")
    lines.append("")
    lines.append("6. 품질 및 무결성 자가진단 결과")
    lines.append(f"   - CR 반영률: {readiness['cr_reflected']}/{readiness['cr_total']} "
                  f"({readiness['cr_rate']*100:.0f}%)")
    lines.append(f"   - 스냅샷 기록: {'있음' if readiness['snapshot_exists'] else '★★ 없음 (등록 필요)'}")
    lines.append("-" * 70)
    lines.append("*본 문서는 DRAFT이며, 연구원의 감수·편집을 거쳐 확정해야 한다. "
                  "파일명(*_draft.md)은 확정 전까지 변경하지 않는다.*")
    return "\n".join(lines)


def inject_demo_hash(draft_text, readiness):
    """실제 빌드 zip이 없는 데모 환경이므로, 대상 행들의 체크섬을 이어붙여
    계산한 대표 해시로 스냅샷 해시 자리를 채운다 (진짜 빌드 스냅샷 해시가 아님을 명시)."""
    combined = "|".join(_sval(row["checksum"]) for row in readiness["target_rows"])
    demo_hash = "sha256:" + hashlib.sha256(combined.encode("utf-8")).hexdigest()
    final_text = draft_text.replace(
        "`<릴리즈 확정 시 자동 주입 예정>`", f"`{demo_hash}` (데모용 대표 해시)"
    )
    final_text = final_text.replace("Document Version: DRAFT (자동 생성, 연구원 검토·확정 전)",
                                     "Document Version: v1.0 (확정, 데모 해시 주입 완료)")
    return final_text, demo_hash


# ---------------------------------------------------------------
# Streamlit UI
# ---------------------------------------------------------------
def render():
    with st.expander("📁 SSoT 데이터 소스", expanded=False):
        uploaded_file = st.file_uploader(
            "또는 보유하신 SSoT 템플릿(.xlsm/.xlsx) 업로드", type=["xlsm", "xlsx"], key="ssot_uploader"
        )

    if uploaded_file is not None:
        target_file = uploaded_file
        display_name = uploaded_file.name
        file_valid = True
    else:
        target_file = _get_valid_file_path(DEFAULT_EXCEL_FILENAME)
        display_name = target_file if target_file else DEFAULT_EXCEL_FILENAME
        file_valid = target_file is not None and os.path.exists(target_file)

    st.title("📦 릴리즈 준비상태 점검 · 릴리즈 노트(Draft)  작성")
    st.caption("SSoT_Template_v3.xlsm ('인스턴스_통합'/'SUP9_이슈현황'/'스냅샷_보관') 기반 순수 Python 데모")

    if not file_valid:
        st.warning(f"경고: `{DEFAULT_EXCEL_FILENAME}` 파일이 `data/` 폴더 또는 메인 폴더에 없습니다. 업로드해 주세요.")
        return

    st.success(f"현재 연동 파일: **{display_name}**")

    try:
        wb = _load_workbook(target_file)
    except Exception as e:
        st.error(f"엑셀 파일 로드 중 오류 발생: {e}")
        return

    missing = [s for s in (SHEET_INSTANCE, SHEET_SUP9, SHEET_SNAPSHOT) if s not in wb.sheetnames]
    if missing:
        st.error(f"필수 시트를 찾을 수 없습니다: {missing} (시트 목록: {wb.sheetnames})")
        return

    instance_rows = _read_instance_rows(wb[SHEET_INSTANCE])
    sup9_rows = _read_sup9(wb[SHEET_SUP9])
    snapshot_baselines = _read_snapshot_baselines(wb[SHEET_SNAPSHOT])

    baselines = _discover_baselines(instance_rows)
    if not baselines:
        st.info("인스턴스_통합 시트에서 baseline_id를 찾지 못했습니다.")
        return

    c1, c2 = st.columns([2, 2])
    with c1:
        baseline_id = st.selectbox("🎯 점검할 Baseline ID 선택", sorted(baselines.keys()))
    with c2:
        report_date = st.date_input("📆 릴리즈 일자", value=date.today())

    meta = baselines[baseline_id]
    st.caption(f"→ 제어기: **{meta.get('controller_id')}** · 적용 차종: **{meta.get('vehicle')}**")

    if st.button("🚀 릴리즈 준비도 점검 실행", type="primary"):
        readiness = compute_readiness(instance_rows, sup9_rows, snapshot_baselines, baseline_id)
        st.session_state["ssot_readiness"] = {
            "baseline_id": baseline_id, "meta": meta, "readiness": readiness,
            "report_date": report_date.isoformat(),
        }
        st.session_state.pop("ssot_draft_note", None)
        st.session_state.pop("ssot_final_note", None)

    saved = st.session_state.get("ssot_readiness")
    if not saved or saved["baseline_id"] != baseline_id:
        return

    readiness = saved["readiness"]

    st.markdown("---")
    st.subheader("📋 릴리즈 준비도 대시보드")
    verdict_color = "🟢" if readiness["verdict"] == "GO" else "🔴"
    m1, m2, m3, m4 = st.columns(4)
    m1.metric("종합판정", f"{verdict_color} {readiness['verdict']}")
    m2.metric("CR 반영률", f"{readiness['cr_reflected']}/{readiness['cr_total']} ({readiness['cr_rate']*100:.0f}%)")
    m3.metric("CVE 미조치/조치중", f"{len(readiness['cve_pending'])}건")
    m4.metric("미종결 Critical 이슈", f"{len(readiness['critical_open'])}건")

    if readiness["reasons"]:
        st.error("**판정 사유(NO-GO):** " + " / ".join(readiness["reasons"]))
    else:
        st.success("모든 점검 항목을 충족했습니다 (GO).")

    st.dataframe(pd.DataFrame(readiness["target_rows"]), use_container_width=True)

    st.markdown("---")
    st.subheader("📝 Draft 릴리즈노트")
    if st.button("✏️ Draft 릴리즈노트 생성"):
        draft_text = generate_draft_release_note(
            baseline_id, saved["meta"], readiness, saved["report_date"]
        )
        st.session_state["ssot_draft_note"] = draft_text
        st.session_state.pop("ssot_final_note", None)

    draft_text = st.session_state.get("ssot_draft_note")
    if draft_text:
        n_warn = draft_text.count("★★")
        if n_warn:
            st.warning(f"⚠️ 미비/미확인 항목 {n_warn}건에 ★★ 표시가 있습니다. 확정 전 반드시 확인·보완하세요.")
        st.code(draft_text, language="text")
        st.download_button(
            "📥 Draft 릴리즈노트 다운로드 (.md)",
            data=draft_text.encode("utf-8"),
            file_name=f"ReleaseNote_{baseline_id}_draft.md",
            mime="text/markdown",
        )

        st.markdown("---")
        st.subheader("✅ 해시 주입 및 최종 확정 (데모)")
        st.caption("실제 빌드 산출물이 없는 데모 환경이므로, 대상 행들의 체크섬을 이어붙인 대표 해시로 대체합니다.")
        if st.button("🔒 데모 해시 주입 및 확정"):
            final_text, demo_hash = inject_demo_hash(draft_text, readiness)
            st.session_state["ssot_final_note"] = (final_text, demo_hash)

        final = st.session_state.get("ssot_final_note")
        if final:
            final_text, demo_hash = final
            st.success(f"데모 해시 주입 완료: `{demo_hash}`")
            st.code(final_text, language="text")
            st.download_button(
                "📥 최종 릴리즈노트 다운로드 (.md)",
                data=final_text.encode("utf-8"),
                file_name=f"ReleaseNote_{baseline_id}.md",
                mime="text/markdown",
            )

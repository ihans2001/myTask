# mitre_monitoring_page.py
# -------------------------------------------------------------
# MITRE CVE 모니터링 데모 (문서번호 참고: monitoringMitre 에이전트)
# 원본: main_dailyMonitoring_Mitre.py / config.py / cve_scraper.py / excel_io.py
#       / macro_runner.py / README.md (dayMonitoring_Template.xlsm)
#
# [Mac Studio 서버 운영을 위한 조정 사항]
#   1) macro_runner.py(win32com으로 AnalyzeCVEs VBA 매크로 호출)는 Windows
#      전용이라 이 데모에서는 완전히 제외했다. 매크로가 필요하면 로컬
#      Windows+Excel 환경에서 별도로 실행해야 한다.
#   2) schedule 기반 상주 프로세스(매일 10:00 자동 실행) 대신, 이 데모는
#      버튼을 눌러 그때그때 1회 실행하는 방식으로 단순화했다(Streamlit 특성상
#      상주 스케줄러는 별도 백그라운드 프로세스로 운영해야 하므로 범위 밖).
#   3) 원본 코드의 검색/상세조회 로직(Playwright + MITRE 공식 API)은 그대로
#      이식했다 — 실제 배포 서버(인터넷 연결 가능)에서는 그대로 동작한다.
#      다만 이 환경(개발 샌드박스)은 외부망 접근이 제한되어 있어 라이브 호출을
#      검증할 수 없으므로, 오프라인에서도 UI/엑셀 반영 흐름을 확인할 수 있도록
#      "데모 모드"(모의 데이터)를 함께 제공한다.
#   4) "keyword를 입력받아야 한다"는 요구사항에 따라, 대장에 등록된 키워드
#      선택뿐 아니라 자유 텍스트로 즉석 키워드를 입력해 조회할 수 있게 했다.
# -------------------------------------------------------------

import os
import io
import re
import random
from datetime import date, datetime

import streamlit as st
import pandas as pd
import openpyxl

DEFAULT_EXCEL_FILENAME = "dayMonitoring_Template.xlsm"
SHEET_KEYWORD = "Latest_ID"
SHEET_MONITORED = "Detected_ID"
SHEET_SUMMARY = "Detected_Summary"

# Latest_ID 시트 컬럼 (config.py와 동일)
COL_KEYWORD_DATE = 3    # C
COL_KEYWORD_TRIAGE = 4  # D
COL_KEYWORD = 5         # E
COL_LAST_CVE = 6        # F
KEYWORD_START_ROW = 5

# Detected_ID 시트 컬럼
COL_MON_NO = 2
COL_MON_DATE = 3
COL_MON_TRIAGE = 4
COL_MON_KEYWORD = 5
COL_MON_CVE_ID = 6
COL_MON_CNA = 7
COL_MON_DESC = 8
MONITORED_START_ROW = 5

# Detected_Summary 시트 컬럼
COL_SUM_TRIAGE = 3
COL_SUM_TRIGGER = 4
COL_SUM_COUNTER = 5
SUMMARY_START_ROW = 5

CVE_API_DETAIL_URL = "https://cveawg.mitre.org/api/cve/{cve_id}"
CVE_SEARCH_URL = "https://www.cve.org/CVERecord/SearchResults?query={keyword}"
_CVE_ID_PATTERN = re.compile(r"CVE-\d{4}-\d{4,7}")
_CVE_PATTERN = re.compile(r"CVE-(\d{4})-(\d+)")


# ---------------------------------------------------------------
# CVE ID 비교 유틸 (cve_scraper.py 그대로 이식)
# ---------------------------------------------------------------
def cve_sort_key(cve_id: str):
    m = _CVE_PATTERN.match(cve_id.strip())
    if not m:
        return (0, 0)
    return (int(m.group(1)), int(m.group(2)))


def is_newer(cve_id: str, baseline_id):
    if not baseline_id:
        return True
    return cve_sort_key(cve_id) > cve_sort_key(baseline_id)


# ---------------------------------------------------------------
# 엑셀 로딩
# ---------------------------------------------------------------
def _get_valid_file_path(filename: str):
    data_folder_path = os.path.join("data", filename)
    if os.path.exists(data_folder_path):
        return data_folder_path
    elif os.path.exists(filename):
        return filename
    return None


def _load_workbook(file_source):
    if isinstance(file_source, str):
        with open(file_source, "rb") as f:
            file_bytes = f.read()
    else:
        file_source.seek(0)
        file_bytes = file_source.read()
    return openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, keep_vba=False)


def read_keywords(wb):
    ws = wb[SHEET_KEYWORD]
    entries = []
    row = KEYWORD_START_ROW
    while True:
        kw = ws.cell(row=row, column=COL_KEYWORD).value
        if kw is None or str(kw).strip() == "":
            break
        entries.append({
            "row": row,
            "keyword": str(kw).strip(),
            "last_cve_id": str(ws.cell(row=row, column=COL_LAST_CVE).value or "").strip() or None,
            "triage": ws.cell(row=row, column=COL_KEYWORD_TRIAGE).value,
            "last_date": ws.cell(row=row, column=COL_KEYWORD_DATE).value,
        })
        row += 1
    return entries


# ---------------------------------------------------------------
# 실시간 조회 (원본 로직 그대로 이식 — 실제 서버 배포 시 인터넷 필요)
# ---------------------------------------------------------------
def fetch_cve_detail_live(cve_id: str) -> dict:
    """MITRE 공식 CVE Services API에서 상세 정보를 조회한다 (requests만 사용, 브라우저 불필요)."""
    import requests

    detail = {"cve_id": cve_id, "cna": "", "description": "", "updated": "", "published": "",
              "cwe": "", "cvss_v3": "", "vendor": "", "product": ""}
    url = CVE_API_DETAIL_URL.format(cve_id=cve_id)
    resp = requests.get(url, timeout=15, headers={"User-Agent": "monitoringMitre-demo/1.0"})
    resp.raise_for_status()
    data = resp.json()

    meta = data.get("cveMetadata", {}) or {}
    cna = (data.get("containers", {}) or {}).get("cna", {}) or {}

    detail["cna"] = meta.get("assignerShortName", "") or ""
    detail["published"] = meta.get("datePublished", "") or ""
    detail["updated"] = meta.get("dateUpdated", "") or ""
    for d in cna.get("descriptions", []) or []:
        if str(d.get("lang", "")).lower().startswith("en"):
            detail["description"] = d.get("value", "") or ""
            break

    vendors, products = [], []
    for item in cna.get("affected", []) or []:
        if item.get("vendor") and item["vendor"] not in vendors:
            vendors.append(item["vendor"])
        if item.get("product") and item["product"] not in products:
            products.append(item["product"])
    detail["vendor"] = "; ".join(vendors)
    detail["product"] = "; ".join(products)

    for m in cna.get("metrics", []) or []:
        for key in ("cvssV3_1", "cvssV3_0"):
            if key in m:
                score = m[key].get("baseScore")
                detail["cvss_v3"] = f"{score}" if score is not None else ""
                break

    return detail


def search_cve_live(keyword: str, baseline_id, max_pages: int = 3):
    """cve.org 검색 결과에서 CVE ID를 최신순으로 추출한다 (Playwright 필요, 실제
    인터넷 연결이 되는 서버에서만 동작한다)."""
    from playwright.sync_api import sync_playwright

    found_ids, seen = [], set()
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        page = browser.new_page()
        try:
            for page_num in range(1, max_pages + 1):
                url = CVE_SEARCH_URL.format(keyword=keyword) + f"&page={page_num}&pageSize=100"
                resp = page.goto(url, timeout=20000)
                if resp is not None and resp.status >= 400:
                    raise RuntimeError(
                        f"cve.org 접속 실패 (HTTP {resp.status}) — 네트워크 방화벽/프록시가 "
                        f"이 사이트 접근을 막고 있을 수 있습니다."
                    )
                page.wait_for_timeout(1500)
                body_text = page.inner_text("body")
                page_ids = []
                for m in _CVE_ID_PATTERN.finditer(body_text):
                    if m.group(0) not in seen:
                        seen.add(m.group(0))
                        page_ids.append(m.group(0))
                if not page_ids:
                    break
                page_ids.sort(key=cve_sort_key, reverse=True)
                reached_baseline = False
                for cid in page_ids:
                    if is_newer(cid, baseline_id):
                        found_ids.append(cid)
                    else:
                        reached_baseline = True
                        break
                if reached_baseline:
                    break
        finally:
            browser.close()
    return found_ids


# ---------------------------------------------------------------
# 데모(오프라인) 모드 — 인터넷 없이 UI/엑셀 반영 흐름을 확인하기 위한 모의 데이터
# ---------------------------------------------------------------
def generate_demo_results(keyword: str, baseline_id):
    """실제 조회 대신, baseline보다 최신인 모의 CVE 2~3건을 생성한다.
    화면과 다운로드 파일 어디에나 '데모(모의)'임을 명시한다."""
    base_year, base_num = cve_sort_key(baseline_id) if baseline_id else (2026, 0)
    rng = random.Random(f"{keyword}-{baseline_id}")
    count = rng.randint(2, 3)
    results = []
    for i in range(count):
        new_num = base_num + rng.randint(50, 900) + i
        cve_id = f"CVE-{base_year}-{new_num}"
        results.append({
            "cve_id": cve_id,
            "cna": "DEMO-CNA",
            "description": f"[데모 데이터] '{keyword}' 키워드 관련 모의 취약점 설명입니다. 실제 조회 결과가 아닙니다.",
            "updated": date.today().isoformat(),
            "published": date.today().isoformat(),
            "cwe": "CWE-000 (데모)",
            "cvss_v3": str(round(rng.uniform(4.0, 9.5), 1)),
            "vendor": "DemoVendor",
            "product": "DemoProduct",
        })
    results.sort(key=lambda r: cve_sort_key(r["cve_id"]), reverse=True)
    return results


# ---------------------------------------------------------------
# 엑셀 반영 (대장에 결과 기록)
# ---------------------------------------------------------------
def apply_results_to_workbook(wb, keyword_entry, keyword, triage, results, today_str):
    ws_kw = wb[SHEET_KEYWORD]
    ws_mon = wb[SHEET_MONITORED]
    ws_sum = wb[SHEET_SUMMARY]

    if keyword_entry is None:
        # 신규 키워드 -> Latest_ID 맨 아래에 새 행 추가
        row = KEYWORD_START_ROW
        while ws_kw.cell(row=row, column=COL_KEYWORD).value not in (None, ""):
            row += 1
        ws_kw.cell(row=row, column=COL_KEYWORD).value = keyword
        ws_kw.cell(row=row, column=COL_KEYWORD_TRIAGE).value = triage or "신규"
        keyword_row = row
    else:
        keyword_row = keyword_entry["row"]

    # Detected_ID 다음 빈 행 찾기
    next_row = MONITORED_START_ROW
    while ws_mon.cell(row=next_row, column=COL_MON_CVE_ID).value not in (None, ""):
        next_row += 1

    for item in results:
        ws_mon.cell(row=next_row, column=COL_MON_DATE).value = today_str
        ws_mon.cell(row=next_row, column=COL_MON_TRIAGE).value = triage
        ws_mon.cell(row=next_row, column=COL_MON_KEYWORD).value = keyword
        ws_mon.cell(row=next_row, column=COL_MON_CVE_ID).value = item["cve_id"]
        ws_mon.cell(row=next_row, column=COL_MON_CNA).value = item.get("cna", "")
        ws_mon.cell(row=next_row, column=COL_MON_DESC).value = item.get("description", "")
        next_row += 1

    if results:
        newest_id = results[0]["cve_id"]
        ws_kw.cell(row=keyword_row, column=COL_LAST_CVE).value = newest_id
        ws_kw.cell(row=keyword_row, column=COL_KEYWORD_DATE).value = today_str

    ws_sum.cell(row=keyword_row, column=COL_SUM_TRIAGE).value = triage
    ws_sum.cell(row=keyword_row, column=COL_SUM_TRIGGER).value = keyword
    ws_sum.cell(row=keyword_row, column=COL_SUM_COUNTER).value = len(results)

    return keyword_row


# ---------------------------------------------------------------
# Streamlit UI
# ---------------------------------------------------------------
def render():
    with st.expander("📁 MITRE 모니터링 데이터 소스", expanded=False):
        uploaded_file = st.file_uploader(
            "또는 보유하신 dayMonitoring 템플릿(.xlsm/.xlsx) 업로드", type=["xlsm", "xlsx"], key="mitre_uploader"
        )

    if uploaded_file is not None:
        target_file = uploaded_file
        display_name = uploaded_file.name
        file_valid = True
    else:
        target_file = _get_valid_file_path(DEFAULT_EXCEL_FILENAME)
        display_name = target_file if target_file else DEFAULT_EXCEL_FILENAME
        file_valid = target_file is not None and os.path.exists(target_file)

    st.title("🛰️ MITRE CVE 모니터링")
    st.caption("main_dailyMonitoring_Mitre.py 등 monitoringMitre 에이전트의 순수 Python 데모 (VBA/Windows 매크로 호출 제외)")

    if not file_valid:
        st.warning(f"경고: `{DEFAULT_EXCEL_FILENAME}` 파일이 `data/` 폴더 또는 메인 폴더에 없습니다. 업로드해 주세요.")
        return

    st.success(f"현재 연동 파일: **{display_name}**")

    try:
        wb = _load_workbook(target_file)
    except Exception as e:
        st.error(f"엑셀 파일 로드 중 오류 발생: {e}")
        return

    if SHEET_KEYWORD not in wb.sheetnames:
        st.error(f"'{SHEET_KEYWORD}' 시트를 찾을 수 없습니다. (시트 목록: {wb.sheetnames})")
        return

    keywords = read_keywords(wb)
    st.subheader("📋 등록된 모니터링 키워드")
    if keywords:
        st.dataframe(pd.DataFrame(keywords)[["row", "keyword", "triage", "last_cve_id", "last_date"]],
                     use_container_width=True)
    else:
        st.info("등록된 키워드가 없습니다. 아래에서 새 키워드로 바로 조회할 수 있습니다.")

    st.markdown("---")
    st.subheader("🔍 키워드 조회")

    mode = st.radio(
        "조회 방식",
        ["🟢 데모 모드 (모의 데이터, 인터넷 불필요)", "🔴 실시간 조회 (Playwright+MITRE API, 인터넷 필요)"],
        horizontal=True,
    )
    is_live = mode.startswith("🔴")

    c1, c2 = st.columns([2, 1])
    with c1:
        existing_kw = ["(직접 입력)"] + [e["keyword"] for e in keywords]
        picked = st.selectbox("등록된 키워드에서 선택 (선택 시 아래 입력창에 채워짐)", existing_kw)
    with c2:
        default_triage = ""
        if picked != "(직접 입력)":
            match = next((e for e in keywords if e["keyword"] == picked), None)
            default_triage = str(match["triage"]) if match and match["triage"] else ""

    keyword_input = st.text_input(
        "🔤 조회할 키워드 입력 (등록되지 않은 새 키워드도 바로 입력 가능)",
        value="" if picked == "(직접 입력)" else picked,
    )
    triage_input = st.text_input("Triage 분류 (선택)", value=default_triage)

    existing_entry = next((e for e in keywords if e["keyword"] == keyword_input.strip()), None)
    baseline_id = existing_entry["last_cve_id"] if existing_entry else None
    if existing_entry:
        st.caption(f"→ 등록된 키워드입니다. 기준(baseline) CVE: `{baseline_id or '없음(최초 조회)'}`")
    elif keyword_input.strip():
        st.caption("→ 대장에 없는 새 키워드입니다. 기준 없이 조회합니다(최초 실행과 동일).")

    if st.button("🚀 조회 실행", type="primary", disabled=not keyword_input.strip()):
        kw = keyword_input.strip()
        try:
            if is_live:
                with st.spinner("cve.org / MITRE API 실시간 조회 중..."):
                    cve_ids = search_cve_live(kw, baseline_id)
                    results = []
                    for cid in cve_ids:
                        try:
                            results.append(fetch_cve_detail_live(cid))
                        except Exception as e:
                            results.append({"cve_id": cid, "cna": "", "description": f"상세 조회 실패: {e}",
                                             "updated": "", "published": "", "cwe": "", "cvss_v3": "",
                                             "vendor": "", "product": ""})
            else:
                results = generate_demo_results(kw, baseline_id)

            st.session_state["mitre_result"] = {
                "keyword": kw, "triage": triage_input.strip(), "results": results,
                "is_live": is_live, "existing_entry": existing_entry,
            }
        except Exception as e:
            st.error(
                f"실시간 조회 중 오류가 발생했습니다: {e}\n\n"
                "이 환경(개발 샌드박스)은 외부망(cve.org, cveawg.mitre.org) 접근이 차단되어 있을 수 있습니다. "
                "실제 서버(맥스튜디오 등, 인터넷 연결됨)에 배포한 뒤 다시 시도하거나, 데모 모드로 UI 흐름을 먼저 확인해 보세요."
            )

    result = st.session_state.get("mitre_result")
    if not result or result["keyword"] != keyword_input.strip():
        return

    st.markdown("---")
    tag = "🔴 실시간 조회 결과" if result["is_live"] else "🟢 데모(모의) 조회 결과 — 실제 취약점 데이터가 아닙니다"
    st.subheader(tag)

    results = result["results"]
    if not results:
        st.info(f"'{result['keyword']}' 키워드에서 기준(baseline)보다 새로운 CVE가 없습니다.")
        return

    st.dataframe(pd.DataFrame(results), use_container_width=True)

    st.markdown("---")
    st.subheader("💾 대장(Excel)에 반영")
    st.caption("Detected_ID/Detected_Summary/Latest_ID 시트에 이번 결과를 기록한 사본을 내려받습니다. "
               "VBA 매크로(AnalyzeCVEs) 자동 호출은 Windows 전용이라 이 데모에서는 생략됩니다.")
    if st.button("📥 대장 반영 후 다운로드용 파일 생성"):
        today_str = date.today().isoformat()
        apply_results_to_workbook(
            wb, result["existing_entry"], result["keyword"], result["triage"], results, today_str
        )
        buf = io.BytesIO()
        wb.save(buf)
        st.session_state["mitre_export"] = buf.getvalue()

    export_bytes = st.session_state.get("mitre_export")
    if export_bytes:
        st.download_button(
            "다운로드: dayMonitoring_updated.xlsx",
            data=export_bytes,
            file_name="dayMonitoring_updated.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
